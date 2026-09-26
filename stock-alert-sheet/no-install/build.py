"""설치 없는 버전: 수식(GOOGLEFINANCE)만으로 동작하는 시트를 xlsx 로 만든다.
구글 드라이브에 올려 구글 스프레드시트로 변환하면 바로 작동한다. (엑셀에서는 GOOGLEFINANCE 가 동작하지 않음)
사용: python3 build.py [출력 경로]
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule

OUT = sys.argv[1] if len(sys.argv) > 1 else '주식_대시보드.xlsx'
ROWS = 50
BOLD = Font(bold=True)
HEAD_FILL = PatternFill('solid', fgColor='DDE3EA')
GREEN = PatternFill('solid', fgColor='C8E6C9')
RED = PatternFill('solid', fgColor='FFCDD2')
YELLOW = PatternFill('solid', fgColor='FFF3C4')

wb = Workbook()

# ── 관심종목 ──
ws = wb.active
ws.title = '관심종목'
headers = ['시장', '종목코드', '종목명', '현재가', '등락률(%)', '목표가', '손절가', '급등락 기준(%)', '상태']
ws.append(headers)
# 목표가·손절가는 시세에 따라 틀린 예시가 되므로 비워 둔다
samples = [('KRX', '005930', '', ''), ('KRX', '000660', '', '')]
for r in range(2, ROWS + 2):
    market, code, target, stop = samples[r - 2] if r - 2 < len(samples) else ('', '', '', '')
    ws.cell(r, 1, market)
    c = ws.cell(r, 2, code)
    c.number_format = '@'
    t = f'A{r}&":"&B{r}'
    ws.cell(r, 3, f'=IF(B{r}="","",IFERROR(GOOGLEFINANCE({t},"name"),"종목 확인불가"))')
    ws.cell(r, 4, f'=IF(B{r}="","",IFERROR(GOOGLEFINANCE({t},"price"),""))').number_format = '#,##0'
    ws.cell(r, 5, f'=IF(B{r}="","",IFERROR(GOOGLEFINANCE({t},"changepct"),""))').number_format = '+0.00;-0.00;0.00'
    ws.cell(r, 6, target).number_format = '#,##0'
    ws.cell(r, 7, stop).number_format = '#,##0'
    ws.cell(r, 8, 5 if r - 2 < len(samples) else '')
    ws.cell(r, 9, (
        f'=IF(D{r}="","",'
        f'IF(AND(F{r}<>"",D{r}>=F{r}),"🎯 목표가 도달",'
        f'IF(AND(G{r}<>"",D{r}<=G{r}),"⚠️ 손절가 이탈",'
        f'IF(AND(H{r}<>"",E{r}>=H{r}),"🚀 급등",'
        f'IF(AND(H{r}<>"",E{r}<=-H{r}),"📉 급락","정상")))))'
    ))
rng = f'A2:I{ROWS + 1}'
ws.conditional_formatting.add(rng, FormulaRule(formula=['LEFT($I2,2)="🎯"'], fill=GREEN))
ws.conditional_formatting.add(rng, FormulaRule(formula=['LEFT($I2,2)="⚠️"'], fill=RED))
ws.conditional_formatting.add(rng, FormulaRule(formula=['OR(LEFT($I2,2)="🚀",LEFT($I2,2)="📉")'], fill=YELLOW))
for col, w in zip('ABCDEFGHI', [8, 10, 18, 11, 10, 11, 11, 13, 16]):
    ws.column_dimensions[col].width = w

# ── 시장지표 ──
mi = wb.create_sheet('시장지표')
mi.append(['분류', '지표', '현재값', '전일 대비(%)', '읽는 법'])
indicators = [
    ('심리', 'VIX (미국 공포지수)', 'INDEXCBOE:VIX', '0.00', '20 이하 안정 · 30 이상 공포'),
    ('환율', '원/달러', 'CURRENCY:USDKRW', '#,##0.0', '급등하면 외국인 매도 압력'),
    ('해외', '미국 S&P500', 'INDEXSP:.INX', '#,##0.00', '미국 증시 방향'),
    ('해외', '미국 나스닥', 'INDEXNASDAQ:.IXIC', '#,##0.00', '기술주 분위기'),
]
for i, (cat, name, ticker, fmt, hint) in enumerate(indicators, start=2):
    mi.cell(i, 1, cat)
    mi.cell(i, 2, name)
    # 환율은 "price" 속성을 붙이면 실패하므로 속성 없이 부른다
    attr = '' if ticker.startswith('CURRENCY:') else ',"price"'
    mi.cell(i, 3, f'=IFERROR(GOOGLEFINANCE("{ticker}"{attr}),"불러오기 실패")').number_format = fmt
    if ticker.startswith('CURRENCY:'):
        mi.cell(i, 4, '')
    else:
        mi.cell(i, 4, f'=IFERROR(GOOGLEFINANCE("{ticker}","changepct"),"")').number_format = '+0.00;-0.00;0.00'
    mi.cell(i, 5, hint)
note_row = len(indicators) + 3
mi.cell(note_row, 1, '※ 구글 시세는 약 20분 늦습니다. 수급·예탁금·금리는 설치형(알림) 버전에서 기록됩니다.')
for col, w in zip('ABCDE', [8, 22, 16, 13, 32]):
    mi.column_dimensions[col].width = w

# ── 사용법 ──
guide = wb.create_sheet('사용법')
for line in [
    '📈 주식 대시보드 (설치 없는 버전)',
    '',
    '1. [관심종목] 시트의 시장·종목코드만 입력하면 종목명·현재가·등락률이 자동으로 채워집니다.',
    '   · 시장: 코스피 = KRX, 코스닥 = KOSDAQ  · 종목코드: 6자리 (예: 005930)',
    '2. 목표가·손절가·급등락 기준(%)을 적으면 상태 칸에 표시되고 줄 색이 바뀝니다.',
    '   · 초록 = 목표가 도달, 빨강 = 손절가 이탈, 노랑 = 급등락',
    '3. [시장지표] 시트에서 VIX·환율·미국 지수를 확인합니다.',
    '',
    '※ 이 버전은 시트를 열었을 때 확인하는 방식입니다. 텔레그램 알림은 오지 않습니다.',
    '※ 투자 판단의 참고용이며, 매매 책임은 사용자에게 있습니다.',
]:
    guide.append([line])
guide['A1'].font = Font(bold=True, size=14)
guide.column_dimensions['A'].width = 90

for sheet in (ws, mi):
    for cell in sheet[1]:
        cell.font = BOLD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal='center')
    sheet.freeze_panes = 'A2'

wb.save(OUT)
print('saved', OUT)
